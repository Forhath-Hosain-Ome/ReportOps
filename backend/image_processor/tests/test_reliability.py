"""Durable lifecycle regressions using real models, paths, UUIDs and API routes."""
import io
import shutil
import subprocess
import tempfile
import uuid
import os
from unittest import skipUnless
from datetime import timedelta
from pathlib import Path
from unittest.mock import Mock, patch

from django.contrib.auth.models import User
from django.db import DatabaseError
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document
from PIL import Image
from rest_framework.test import APIClient
from rest_framework.exceptions import ValidationError

from image_processor import lifecycle
from image_processor.coordination import StaleExecution, execution_write, batch_lock
from image_processor.models import FolderBatch, FolderReport, FolderImage, ProcessingAttempt, ImageOutcome, OperationalSettings
from image_processor.output_paths import OutputPathError, stored_output_path
from image_processor.retention_paths import source_directory, output_directory_for_batch
from image_processor.tasks import process_folder_task
from image_processor.tasks.cleanup_task import cleanup_temp_dirs_task


class ReliabilityFixture(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name).resolve()
        self.staging = self.base / "batch_uploads"
        self.staging.mkdir()
        self.output = self.base / "media" / "output"
        self.template = self.base / "template.docx"
        doc = Document()
        doc.add_table(rows=1, cols=1)
        doc.save(self.template)
        self.settings = override_settings(
            MEDIA_ROOT=self.base / "media", IMAGE_PROCESSOR_STAGING_ROOT=self.staging,
            IMAGE_PROCESSOR_SETTINGS={"OUTPUT_DIR": self.output, "DEFECT_IMAGE_TEMPLATE_PATH": self.template},
        )
        self.settings.enable()
        self.addCleanup(self.settings.disable)
        self.user = User.objects.create_user(username="owner")
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def make_batch(self, files=None):
        files = files or {"A/one.png": True}
        root = self.staging / str(uuid.uuid4())
        root.mkdir()
        manifest = []
        for relative, valid in files.items():
            path = root / relative
            path.parent.mkdir(exist_ok=True, parents=True)
            if valid:
                Image.new("RGB", (8, 8), "red").save(path)
            else:
                path.write_bytes(b"not an image")
            manifest.append({"relative_path": relative, "original_filename": path.name})
        batch = lifecycle.create_intake(self.user, root, manifest, {"date": "2024-02-29"})
        self.publish(batch)
        batch.refresh_from_db()
        return batch

    def publish(self, batch):
        task = Mock()
        task.delay.return_value.id = "test-task"
        return lifecycle.publish_dispatch(batch.pk, batch.dispatch_generation, batch.dispatch_token, task)

    def run_worker(self, batch, **kwargs):
        with patch("image_processor.tasks.process_image._push"):
            return process_folder_task.run(batch.pk, batch.source_folder,
                dispatch_generation=batch.dispatch_generation, dispatch_token=str(batch.dispatch_token), **kwargs)

    def retry(self, batch, ids=None):
        with patch("image_processor.lifecycle.publish_dispatch"):
            new = lifecycle.retry_batch(batch.pk, self.user, ids)
        self.publish(new)
        new.refresh_from_db()
        return new


class LifecycleTests(ReliabilityFixture):
    def test_source_permission_failure_is_not_an_image_failure(self):
        import errno
        batch = self.make_batch()
        with patch("image_processor.tasks.process_image.Image.open", side_effect=PermissionError(errno.EACCES, "unreadable")), self.assertRaises(PermissionError):
            self.run_worker(batch)
        report = batch.reports.get()
        self.assertEqual(report.status, "FAILED")
        self.assertEqual(report.failed_images, 0)
        self.assertFalse(ImageOutcome.objects.exists())

    def test_retry_progress_counts_preserved_terminal_work_only(self):
        batch = self.make_batch({"A/good.png": True, "B/bad.png": False})
        self.run_worker(batch)
        retry = self.retry(batch, [batch.reports.get(folder_name="B").pk])
        self.assertEqual(retry.progress_percent, 50)
        self.assertEqual(retry.failed_folders, 0)

    def test_read_revokes_missing_artifact_without_mutating_attempt(self):
        batch = self.make_batch()
        self.run_worker(batch)
        report = batch.reports.get()
        attempt = report.current_attempt
        lifecycle.authoritative_path(report).unlink()
        response = self.client.get(f"/image/folder/reports/{report.pk}/")
        self.assertEqual(response.data["status"], "FAILED")
        self.assertIsNone(response.data["current_attempt_id"])
        attempt.refresh_from_db()
        self.assertEqual(attempt.status, "COMPLETED")

    def test_stale_recovery_preserves_valid_partial_and_completed_results(self):
        batch = self.make_batch({"A/good.png": True, "A/bad.png": False, "B/good.png": True, "C/good.png": True})
        original_finish = lifecycle.finish_report
        def interrupted(owned, attempt_id, results, output):
            original_finish(owned, attempt_id, results, output)
            if owned.reports.filter(status__in=["COMPLETED", "PARTIAL"]).count() == 2:
                raise StaleExecution("simulated process loss")
        with patch("image_processor.lifecycle.finish_report", side_effect=interrupted), self.assertRaises(StaleExecution):
            self.run_worker(batch)
        FolderBatch.objects.filter(pk=batch.pk).update(lease_expires_at=timezone.now() - timedelta(seconds=1))
        valid = list(batch.reports.filter(status__in=["COMPLETED", "PARTIAL"]).values_list("pk", "current_attempt_id"))
        retry = self.retry(batch)
        self.assertEqual(len(retry.intended_reports), 1)
        self.run_worker(retry)
        for report_id, attempt_id in valid:
            self.assertEqual(FolderReport.objects.get(pk=report_id).current_attempt_id, attempt_id)

    def test_manifest_mismatch_cannot_complete_batch(self):
        batch = self.make_batch()
        original_reconcile = lifecycle.reconcile
        def mismatch(owned):
            FolderReport.objects.create(batch=owned, folder_name="unexpected")
            return original_reconcile(owned)
        with patch("image_processor.lifecycle.reconcile", side_effect=mismatch), self.assertRaises(RuntimeError):
            self.run_worker(batch)
        batch.refresh_from_db()
        self.assertEqual(batch.last_error_code, "BATCH_RECONCILIATION_FAILED")

    def test_publisher_unknown_never_overwrites_queued_state(self):
        batch = self.make_batch()
        FolderBatch.objects.filter(pk=batch.pk).update(dispatch_state="PUBLISHING")
        lifecycle._publication_result(batch.pk, batch.dispatch_generation, batch.dispatch_token, "QUEUED", "")
        lifecycle._publication_result(batch.pk, batch.dispatch_generation, batch.dispatch_token, "UNKNOWN", "DISPATCH_OUTCOME_UNKNOWN")
        batch.refresh_from_db()
        self.assertEqual(batch.dispatch_state, "QUEUED")

    def test_nested_and_loose_images_have_independent_manifest_reports(self):
        batch = self.make_batch({"root.png": True, "A/one.png": True, "A/nested/two.png": True})
        self.assertEqual(set(batch.reports.values_list("folder_name", flat=True)), {".", "A", "A/nested"})
        self.assertEqual(FolderImage.objects.count(), 3)
        self.run_worker(batch)
        batch.refresh_from_db()
        self.assertEqual(batch.status, "COMPLETED")
        self.assertEqual(batch.processed_folders, 3)
        self.assertEqual(batch.successful_images + batch.failed_images, batch.submitted_images)
        for report in batch.reports.all():
            self.assertEqual(report.successful_images + report.failed_images, report.submitted_images)
            self.assertEqual(len(Document(lifecycle.authoritative_path(report)).inline_shapes), 1)
        self.assertTrue(Path(batch.source_folder, "root.png").is_file())

    def test_partial_and_all_image_failures_have_truthful_accounting(self):
        for good in (True, False):
            batch = self.make_batch({"A/good.png": good, "A/bad.png": False})
            self.run_worker(batch)
            report = batch.reports.get()
            batch.refresh_from_db()
            self.assertEqual(report.status, "PARTIAL" if good else "FAILED")
            self.assertEqual(report.successful_images + report.failed_images, 2)
            self.assertEqual(batch.processed_folders + batch.partial_reports + batch.failed_folders, 1)
            self.assertEqual(batch.successful_images + batch.failed_images, 2)
            self.assertEqual(batch.progress_percent, 100)
            self.assertEqual(bool(report.current_attempt_id), good)
            self.assertEqual(report.attempts.get().outcomes.count(), 2)

    def test_duplicate_delivery_and_missing_identity_do_not_mutate(self):
        batch = self.make_batch()
        self.run_worker(batch)
        counts = (FolderReport.objects.count(), ProcessingAttempt.objects.count(), ImageOutcome.objects.count())
        self.assertIn("ignored", self.run_worker(batch))
        self.assertIn("ignored", process_folder_task.run(batch.pk, batch.source_folder))
        self.assertEqual(counts, (FolderReport.objects.count(), ProcessingAttempt.objects.count(), ImageOutcome.objects.count()))

    def test_selective_retry_retains_other_report_and_attempt_history(self):
        batch = self.make_batch({"A/one.png": True, "B/bad.png": False})
        self.run_worker(batch)
        old = batch.reports.get(folder_name="A")
        failed = batch.reports.get(folder_name="B")
        previous = old.current_attempt_id
        Image.new("RGB", (8, 8)).save(Path(batch.source_folder) / "B/bad.png")
        batch = self.retry(batch, [failed.pk])
        self.run_worker(batch)
        old.refresh_from_db()
        failed.refresh_from_db()
        self.assertEqual(old.current_attempt_id, previous)
        self.assertEqual(old.attempts.count(), 1)
        self.assertEqual(failed.attempts.count(), 2)
        self.assertEqual(failed.status, "COMPLETED")

    def test_preserved_artifact_disappearance_clears_authority_only(self):
        batch = self.make_batch({"A/one.png": True, "B/bad.png": False})
        self.run_worker(batch)
        report = batch.reports.get(folder_name="A")
        attempt = report.current_attempt
        path = lifecycle.authoritative_path(report)
        batch = self.retry(batch, [batch.reports.get(folder_name="B").pk])
        path.unlink()
        self.run_worker(batch)
        report.refresh_from_db()
        attempt.refresh_from_db()
        self.assertEqual(report.status, "FAILED")
        self.assertIsNone(report.current_attempt_id)
        self.assertEqual(report.pdf_output_path, "")
        self.assertEqual(attempt.status, "COMPLETED")
        self.assertTrue(attempt.output_path)

    def test_current_dispatch_artifact_disappearance_before_reconciliation(self):
        batch = self.make_batch()
        reconcile = lifecycle.reconcile
        def remove_and_reconcile(owned):
            report = owned.reports.get()
            lifecycle.authoritative_path(report).unlink()
            return reconcile(owned)
        with patch("image_processor.lifecycle.reconcile", side_effect=remove_and_reconcile):
            self.run_worker(batch)
        self.assertEqual(batch.reports.get().status, "FAILED")
        self.assertEqual(batch.reports.get().attempts.get().status, "COMPLETED")

    def test_atomic_authority_rollback_does_not_leave_success_outcomes(self):
        batch = self.make_batch({"A/one.png": True, "A/two.png": True})
        save = FolderReport.save
        def fail_authority(report, *args, **kwargs):
            if report.current_attempt_id:
                raise DatabaseError("simulated authority failure")
            return save(report, *args, **kwargs)
        with patch.object(FolderReport, "save", fail_authority), self.assertRaises(DatabaseError):
            self.run_worker(batch)
        report = batch.reports.get()
        self.assertEqual(report.status, "FAILED")
        self.assertIsNone(report.current_attempt_id)
        self.assertFalse(report.pdf_output_path)
        self.assertEqual(ImageOutcome.objects.filter(status="SUCCEEDED").count(), 0)
        self.assertEqual(report.failed_images, 0)  # No fabricated image failures.

    def test_initialization_failure_terminalizes_and_reraises(self):
        batch = self.make_batch()
        self.template.unlink()
        with self.assertRaises(FileNotFoundError):
            self.run_worker(batch)
        batch.refresh_from_db()
        self.assertEqual(batch.status, "FAILED")
        self.assertIsNone(batch.active_run_token)
        self.assertEqual(batch.reports.get().attempts.get().error_code, "EXECUTION_FAILED")
        self.assertTrue(Path(batch.source_folder).exists())

    def test_irreconcilable_batch_has_explicit_terminal_failure(self):
        batch = self.make_batch()
        with patch("image_processor.lifecycle.reconcile", side_effect=RuntimeError("reconcile")), self.assertRaises(RuntimeError):
            self.run_worker(batch)
        batch.refresh_from_db()
        self.assertEqual(batch.status, "FAILED")
        self.assertEqual(batch.last_error_code, "BATCH_RECONCILIATION_FAILED")
        self.assertIsNotNone(batch.processing_finished_at)
        self.assertIsNone(batch.active_run_token)
        self.assertIsNone(batch.lease_expires_at)

    def test_stale_tokens_fence_every_write_boundary(self):
        batch = self.make_batch()
        owned = lifecycle.claim_execution(batch.pk, batch.dispatch_generation, batch.dispatch_token)
        attempt = owned.reports.get().attempts.get()
        for field, value in (("active_run_token", uuid.uuid4()), ("dispatch_token", uuid.uuid4()),
                             ("dispatch_generation", 999)):
            original = getattr(owned, field)
            FolderBatch.objects.filter(pk=batch.pk).update(**{field: value})
            for mutation in (lambda: lifecycle.finish_report(owned, attempt.pk, []),
                             lambda: lifecycle.fail_execution(owned), lambda: lifecycle.reconcile(owned),
                             lambda: lifecycle.reconciliation_failed(owned)):
                with self.assertRaises(StaleExecution):
                    mutation()
            FolderBatch.objects.filter(pk=batch.pk).update(**{field: original})
        self.assertEqual(ImageOutcome.objects.count(), 0)

    def test_stale_recovery_covers_requested_queued_processing_intent(self):
        batch = self.make_batch({"A/one.png": True, "B/two.png": True, "C/three.png": True})
        owned = lifecycle.claim_execution(batch.pk, batch.dispatch_generation, batch.dispatch_token)
        for attempt, state in zip(ProcessingAttempt.objects.all(), ["REQUESTED", "QUEUED", "PROCESSING"]):
            attempt.status = state
            attempt.save(update_fields=["status"])
        FolderBatch.objects.filter(pk=batch.pk).update(lease_expires_at=timezone.now() - timedelta(seconds=1))
        recovered = self.retry(batch)
        self.assertEqual(recovered.dispatch_generation, 2)
        self.assertEqual(ProcessingAttempt.objects.filter(generation=2).count(), 3)
        with self.assertRaises(StaleExecution):
            lifecycle.fail_execution(owned)
        self.run_worker(recovered)
        recovered.refresh_from_db()
        self.assertEqual(recovered.status, "COMPLETED")

    def test_active_lease_recovery_and_shared_lock_are_rejected(self):
        batch = self.make_batch()
        lifecycle.claim_execution(batch.pk, batch.dispatch_generation, batch.dispatch_token)
        with self.assertRaisesRegex(lifecycle.LifecycleError, "lease"):
            self.retry(batch)
        with batch_lock(batch.pk):
            with self.assertRaisesRegex(lifecycle.LifecycleError, "processing"):
                self.retry(batch)

    def test_missing_or_expired_source_requires_reupload(self):
        for expired in (False, True):
            batch = self.make_batch({"A/bad.png": False})
            self.run_worker(batch)
            if expired:
                FolderBatch.objects.filter(pk=batch.pk).update(expires_at=timezone.now() - timedelta(seconds=1))
            else:
                shutil.rmtree(batch.source_folder)
            response = self.client.post(f"/image/folder/batches/{batch.pk}/retry/", {}, format="json")
            self.assertEqual(response.status_code, 410)
            self.assertEqual(response.data["code"], "REUPLOAD_REQUIRED")

    def test_dispatch_preflight_failure_and_unknown_are_recoverable(self):
        for preflight in (True, False):
            batch = self.make_batch()
            FolderBatch.objects.filter(pk=batch.pk).update(dispatch_state="REQUESTED")
            task = Mock()
            task.delay.side_effect = ConnectionError("broker")
            with patch("image_processor.lifecycle.dumps", side_effect=ValueError() if preflight else None):
                state = lifecycle.publish_dispatch(batch.pk, batch.dispatch_generation, batch.dispatch_token, task)
            self.assertEqual(state, "FAILED" if preflight else "UNKNOWN")
            self.assertIsNone(lifecycle.claim_execution(batch.pk, batch.dispatch_generation, batch.dispatch_token))
            self.assertEqual(self.retry(batch).dispatch_generation, 2)

    def test_publisher_uncertainty_never_overwrites_advanced_worker(self):
        batch = self.make_batch()
        FolderBatch.objects.filter(pk=batch.pk).update(dispatch_state="REQUESTED")
        def claim_then_fail(*args, **kwargs):
            self.assertIsNotNone(lifecycle.claim_execution(batch.pk, batch.dispatch_generation, batch.dispatch_token))
            raise ConnectionError("publisher lost acknowledgement")
        task = Mock()
        task.delay.side_effect = claim_then_fail
        lifecycle.publish_dispatch(batch.pk, batch.dispatch_generation, batch.dispatch_token, task)
        batch.refresh_from_db()
        self.assertEqual(batch.dispatch_state, "PROCESSING")

    def test_settings_permissions_validation_and_new_job_snapshot(self):
        url = "/image/settings/"
        self.assertEqual(APIClient().patch(url, {"retention_hours": 2}).status_code, 401)
        self.assertEqual(self.client.patch(url, {"retention_hours": 2}).status_code, 403)
        self.user.is_superuser = True
        self.user.save()
        self.assertEqual(self.client.patch(url, {"image_width_px": 0}, format="json").status_code, 400)
        first = self.make_batch()
        self.assertEqual(self.client.patch(url, {"retention_hours": 2, "image_width_px": 80, "image_height_px": 60}, format="json").status_code, 200)
        second = self.make_batch()
        self.assertEqual(first.settings_snapshot["image_width_px"], 325)
        self.assertEqual(second.settings_snapshot["retention_hours"], 2)
        self.run_worker(second)
        prepared = next(self.output.rglob("*.jpg"))
        with Image.open(prepared) as image:
            self.assertEqual(image.size, (80, 60))

    def test_owner_scoped_raw_docx_and_paginated_logs(self):
        batch = self.make_batch()
        self.run_worker(batch)
        report = batch.reports.get()
        url = f"/image/folder/reports/{report.pk}/docx/"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("wordprocessingml", response["Content-Type"])
        payload = b"".join(response.streaming_content)
        self.assertEqual(len(Document(io.BytesIO(payload)).inline_shapes), 1)
        logs = self.client.get(f"/image/folder/batches/{batch.pk}/logs/")
        self.assertIn("results", logs.data)
        self.assertEqual(logs.data["results"][0]["outcomes"][0]["original_filename"], "one.png")
        self.client.force_authenticate(User.objects.create_user(username="other"))
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.post(f"/image/folder/batches/{batch.pk}/retry/", {}).status_code, 404)

    def test_docx_open_race_returns_safe_unavailable_response(self):
        batch = self.make_batch()
        self.run_worker(batch)
        with patch("image_processor.views.folder_docx_download_view.open", side_effect=FileNotFoundError(), create=True):
            response = self.client.get(f"/image/folder/reports/{batch.reports.get().pk}/docx/")
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.data["code"], "ARTIFACT_UNAVAILABLE")

    def test_stored_path_contract_rejects_absolute_traversal_and_links(self):
        for value in ("", "../outside.docx", str(self.template)):
            with self.assertRaises(OutputPathError):
                stored_output_path(value)


class RetentionTests(ReliabilityFixture):
    def test_output_root_and_batch_symlink_to_sibling_are_never_deleted(self):
        batch = self.make_batch()
        self.expire(batch)
        with override_settings(IMAGE_PROCESSOR_SETTINGS={"OUTPUT_DIR": self.output.parent}):
            with self.assertRaises(ValueError):
                output_directory_for_batch(batch)
        sibling = self.staging / str(uuid.uuid4())
        sibling.mkdir()
        original = Path(batch.source_folder)
        shutil.rmtree(original)
        try:
            original.symlink_to(sibling, target_is_directory=True)
        except OSError as exc:
            self.skipTest(str(exc))
        result = cleanup_temp_dirs_task()
        batch.refresh_from_db()
        self.assertGreaterEqual(result["errors"], 1)
        self.assertIsNone(batch.source_deleted_at)
        self.assertTrue(sibling.exists())
        self.assertIsNone(batch.maintenance_token)

    @skipUnless(os.name == "nt", "Windows short path alias regression")
    def test_windows_short_root_alias_resolves_to_exact_owned_directory(self):
        import ctypes
        buffer = ctypes.create_unicode_buffer(32768)
        length = ctypes.windll.kernel32.GetShortPathNameW(str(self.staging), buffer, len(buffer))
        self.assertTrue(length)
        batch = self.make_batch()
        with override_settings(IMAGE_PROCESSOR_STAGING_ROOT=buffer.value):
            self.assertEqual(source_directory(batch), Path(batch.source_folder))

    def expire(self, batch):
        self.run_worker(batch)
        FolderBatch.objects.filter(pk=batch.pk).update(expires_at=timezone.now() - timedelta(hours=1))
        batch.refresh_from_db()

    def test_exact_directories_and_truthful_absence_markers(self):
        batch = self.make_batch()
        self.expire(batch)
        sibling = self.staging / str(uuid.uuid4())
        sibling.mkdir()
        counts = cleanup_temp_dirs_task()
        batch.refresh_from_db()
        self.assertEqual(counts["removed"], 2)
        self.assertTrue(sibling.exists())
        self.assertIsNotNone(batch.source_deleted_at)
        self.assertIsNotNone(batch.output_deleted_at)
        self.assertIsNone(batch.maintenance_token)

    def test_empty_root_outside_sibling_and_traversal_are_refused(self):
        batch = self.make_batch()
        self.expire(batch)
        for invalid in ("", str(self.staging), str(self.base), str(self.staging / str(uuid.uuid4())), str(self.staging / ".." / "outside")):
            FolderBatch.objects.filter(pk=batch.pk).update(source_folder=invalid)
            batch.refresh_from_db()
            with self.assertRaises(ValueError):
                source_directory(batch)
        self.assertTrue(self.staging.exists())

    def test_source_symlink_and_batch_output_redirect_are_refused(self):
        batch = self.make_batch()
        self.expire(batch)
        outside = self.base / "outside"
        outside.mkdir()
        for source in (True, False):
            target = Path(batch.source_folder) if source else output_directory_for_batch(batch)
            shutil.rmtree(target)
            try:
                target.symlink_to(outside, target_is_directory=True)
            except OSError as exc:
                self.skipTest(str(exc))
            with self.assertRaises((ValueError, OutputPathError, ValidationError)):
                (source_directory if source else output_directory_for_batch)(batch)
            target.unlink()
        self.assertTrue(outside.exists())

    def test_deletion_failure_releases_maintenance_and_does_not_count_removed(self):
        batch = self.make_batch()
        self.expire(batch)
        with patch("image_processor.tasks.cleanup_task.shutil.rmtree", side_effect=PermissionError()):
            counts = cleanup_temp_dirs_task()
        batch.refresh_from_db()
        self.assertEqual(counts["removed"], 0)
        self.assertEqual(counts["errors"], 2)
        self.assertIsNone(batch.maintenance_token)
        self.assertIsNone(batch.source_deleted_at)
        self.assertIsNone(batch.output_deleted_at)

    def test_absent_source_does_not_delete_anything(self):
        batch = self.make_batch()
        self.expire(batch)
        shutil.rmtree(batch.source_folder)
        counts = cleanup_temp_dirs_task()
        batch.refresh_from_db()
        self.assertEqual(counts["absent"], 1)
        self.assertIsNotNone(batch.source_deleted_at)


class PdfTests(ReliabilityFixture):
    def test_response_construction_failure_closes_file_and_removes_tree(self):
        from pypdf import PdfWriter
        batch = self.make_batch()
        self.run_worker(batch)
        root = self.base / "conversion"
        root.mkdir()
        def convert(source, destination):
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            with open(destination, "wb") as stream:
                writer.write(stream)
        with patch("image_processor.views.folder_pdf_download_view.tempfile.mkdtemp", return_value=str(root)), patch("image_processor.views.folder_pdf_download_view.convert_docx_to_pdf", side_effect=convert), patch("image_processor.views.folder_pdf_download_view.FileResponse", side_effect=RuntimeError("response")):
            response = self.client.get(f"/image/folder/reports/{batch.reports.get().pk}/pdf/")
        self.assertEqual(response.status_code, 503)
        self.assertFalse(root.exists())

    def test_converter_rejects_a_header_without_a_valid_pdf_structure(self):
        batch = self.make_batch()
        self.run_worker(batch)
        def invalid(source, destination):
            Path(destination).write_bytes(b"%PDF-1.7\nnot a document")
        with patch("image_processor.views.folder_pdf_download_view.convert_docx_to_pdf", side_effect=invalid):
            response = self.client.get(f"/image/folder/reports/{batch.reports.get().pk}/pdf/")
        self.assertEqual(response.status_code, 503)

    def test_timeout_nonzero_and_missing_output_remove_entire_temp_tree(self):
        batch = self.make_batch()
        self.run_worker(batch)
        report = batch.reports.get()
        from image_processor.views import folder_pdf_download_view as pdf
        for effect in (subprocess.TimeoutExpired("libreoffice", 120), subprocess.CalledProcessError(1, "libreoffice"), None):
            root = self.base / str(uuid.uuid4())
            root.mkdir()
            (root / "auxiliary").mkdir()
            with patch.object(pdf.tempfile, "mkdtemp", return_value=str(root)), patch.object(pdf.subprocess, "run", side_effect=effect, return_value=Mock(returncode=0)):
                response = self.client.get(f"/image/folder/reports/{report.pk}/pdf/")
            self.assertEqual(response.status_code, 503)
            self.assertFalse(root.exists())

    def test_successful_pdf_tree_survives_until_response_close(self):
        batch = self.make_batch()
        self.run_worker(batch)
        report = batch.reports.get()
        root = self.base / "conversion"
        root.mkdir()
        def convert(source, destination):
            from pypdf import PdfWriter
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            with open(destination, "wb") as stream:
                writer.write(stream)
            (Path(destination).parent / "profile").mkdir()
        with patch("image_processor.views.folder_pdf_download_view.tempfile.mkdtemp", return_value=str(root)), patch("image_processor.views.folder_pdf_download_view.convert_docx_to_pdf", side_effect=convert):
            response = self.client.get(f"/image/folder/reports/{report.pk}/pdf/")
        self.assertTrue(root.exists())
        self.assertTrue(b"".join(response.streaming_content).startswith(b"%PDF-"))
        self.assertFalse(root.exists())
